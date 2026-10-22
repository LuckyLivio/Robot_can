"""Build the detailed design from reviewed sections using the workspace document runtime."""
from __future__ import annotations
import argparse
from collections import defaultdict
import copy
import hashlib
import json
import math
import re
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from lxml import etree
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
RUN = ROOT / "reproduction/graduation_repairs_20261022_01"
ASSETS = HERE / "assets"
TITLE = "机器人抓取搬运与失败恢复系统详细设计"
MATH_DEPS = Path.home() / "AppData/Local/Temp/robot_can_design_20261022_01/doc_math_deps"
sys.path.insert(0, str(MATH_DEPS))
from latex2mathml.converter import convert as latex_to_mathml


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def font(size, bold=False):
    return ImageFont.truetype("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc", size)


def canvas(height=1080):
    image = Image.new("RGB", (2000, height), "white")
    return image, ImageDraw.Draw(image)


def centered(draw, text, center, size=38, bold=False, fill="#17232a"):
    f = font(size, bold)
    lines = text.split("\n")
    for index, line in enumerate(lines):
        bbox = draw.textbbox((0, 0), line, font=f)
        width = bbox[2] - bbox[0]
        y = center[1] + (index - (len(lines)-1)/2) * (size+13)
        draw.text((center[0]-width/2, y-size/2-5), line, font=f, fill=fill)


def box(draw, xy, title, subtitle="", fill="#edf2f5"):
    draw.rounded_rectangle(xy, radius=14, fill=fill, outline="#5d707c", width=3)
    x, y = (xy[0]+xy[2])/2, (xy[1]+xy[3])/2
    centered(draw, title, (x, y-22 if subtitle else y), size=38, bold=True)
    if subtitle:
        centered(draw, subtitle, (x, y+37), size=30)


def arrow(draw, points, color="#455d6b", width=5):
    draw.line(points, fill=color, width=width, joint="curve")
    a, b = points[-2], points[-1]
    theta = math.atan2(b[1]-a[1], b[0]-a[0])
    end = (b[0], b[1])
    tips = [(b[0]-20*math.cos(theta-.47), b[1]-20*math.sin(theta-.47)),
            (b[0]-20*math.cos(theta+.47), b[1]-20*math.sin(theta+.47))]
    draw.polygon([end, *tips], fill=color)


def save_diagrams():
    ASSETS.mkdir(exist_ok=True)
    im, d = canvas(1080)
    box(d, (80,70,1920,230), "数据与仿真层", "HDF5示范与物理起点     Panda与PickPlaceCan     MuJoCo状态推进")
    for xy, title, sub in [((80,320,620,490),"统一观测","23维特权状态与语义版本"),
                           ((730,320,1270,490),"行为克隆策略","8步动作块与4个候选"),
                           ((1380,320,1920,490),"世界模型筛选","真实队列前缀与保守回退")]:
        box(d,xy,title,sub)
    arrow(d,[(350,230),(350,320)])
    arrow(d,[(620,405),(730,405)])
    arrow(d,[(1270,405),(1380,405)])
    box(d,(730,600,1270,770),"v5恢复监督器","可观测抓持与阶段状态机")
    arrow(d,[(1650,490),(1650,550),(1330,550),(1330,650),(1270,650)])
    arrow(d,[(350,490),(350,685),(730,685)])
    box(d,(1380,600,1920,770),"D2执行队列","出队后加噪并限幅")
    arrow(d,[(1270,730),(1380,730)])
    arrow(d,[(1860,600),(1860,490)],color="#a27436",width=4)
    centered(d,"Q",(1900,545),28,fill="#a27436")
    arrow(d,[(1920,655),(1970,655),(1970,150),(1920,150)])
    box(d,(80,870,1920,1010),"实验与证据层","输入锁   源码与权重哈希   逐步轨迹   起点级配对分析   物理回放",fill="#f5f5f5")
    centered(d,"橙色Q为在途反馈；成功标志只用于评价；恢复不读取未来奖励或模型评分",(1000,825),30)
    im.save(ASSETS/"architecture.png",dpi=(300,300))

    im,d=canvas(860)
    names=[("设计与配置","语义 预算 条件"),("训练或装载","明确检查点规则"),("冻结与新起点","输入锁与重生成"),
           ("执行与留痕","完整控制链"),("审计与物理回放","重跑单独记账"),("配对统计与交付","独立单位为起点")]
    coords=[]
    for index,(title,sub) in enumerate(names):
        row,col=divmod(index,3)
        if row: col=2-col
        xy=(100+col*640,120+row*370,620+col*640,320+row*370)
        coords.append(xy);box(d,xy,title,sub)
    for index in range(5):
        a,b=coords[index:index+2]
        if index==2: arrow(d,[((a[0]+a[2])/2,a[3]),((b[0]+b[2])/2,b[1])])
        elif index<2: arrow(d,[(a[2],(a[1]+a[3])/2),(b[0],(b[1]+b[3])/2)])
        else: arrow(d,[(a[0],(a[1]+a[3])/2),(b[2],(b[1]+b[3])/2)])
    centered(d,"同一物理起点上的方法、噪声与固定种子重复执行，不增加独立起点数",(1000,775),32)
    im.save(ASSETS/"evidence_pipeline.png",dpi=(300,300))

    im,d=canvas(830)
    box(d,(90,90,640,250),"离线示范","只读train与dev")
    box(d,(90,420,640,580),"在线观测","仿真器当前绝对位姿")
    box(d,(850,210,1510,450),"canonical共享构造","绝对姿态归一化\n相对位姿同刻重算",fill="#e5edf2")
    arrow(d,[(640,170),(735,170),(735,285),(850,285)])
    arrow(d,[(640,500),(735,500),(735,375),(850,375)])
    box(d,(850,580,1510,740),"23维几何一致状态","shape 有限值 两绝对姿态范数")
    arrow(d,[(1180,450),(1180,580)])
    box(d,(1640,300,1930,550),"模型契约","语义匹配\n各自标准化")
    arrow(d,[(1510,660),(1785,660),(1785,550)])
    centered(d,"legacy保留旧拼接与旧权重；不会自动获得相同在线几何校验",(1000,795),30)
    im.save(ASSETS/"observation_contract.png",dpi=(300,300))

    im,d=canvas(850)
    cols=[550,1080,1610]
    centered(d,"同一次BC推理输出5个完整动作块分量",(1000,75),40,True)
    box(d,(100,185,620,365),"当前状态23维","训练均值与尺度")
    box(d,(740,185,1260,365),"三层全连接","512 → 512 → 512")
    box(d,(1380,185,1900,365),"混合密度输出头","5 × (1 + 2 × 8 × 7)")
    arrow(d,[(620,275),(740,275)]);arrow(d,[(1260,275),(1380,275)])
    box(d,(100,515,620,705),"概率权重","5个混合logits")
    box(d,(740,515,1260,705),"动作块均值","5 × 8 × 7，tanh")
    box(d,(1380,515,1900,705),"对角标准差","softplus与限幅")
    for x in (360,1000,1640):
        arrow(d,[(1640,365),(1640,430),(x,430),(x,515)])
    centered(d,"训练使用完整动作块负对数似然；部署选最高权重分量的均值块",(1000,785),32)
    im.save(ASSETS/"policy_network.png",dpi=(300,300))

    im,d=canvas(800)
    centered(d,"D2动作执行时序",(1000,60),42,True)
    x0,width=150,560
    for index,label in enumerate(("控制时刻 t","控制时刻 t+1","控制时刻 t+2")):
        box(d,(x0+index*width,140,x0+index*width+480,280),label)
    lines=[("新提交",["u(t)","u(t+1)","u(t+2)"]),
           ("加噪前执行",["旧队列 q0","旧队列 q1","u(t)"]),
           ("执行后队列",["[q1, u(t)]","[u(t), u(t+1)]","[u(t+1), u(t+2)]"])]
    for row,(label,values) in enumerate(lines):
        y=370+row*140
        centered(d,label,(180,y),32,True)
        for index,value in enumerate(values): centered(d,value,(640+index*510,y),34)
        d.line((60,y+65,1940,y+65),fill="#c5cdd2",width=2)
    centered(d,"实际动作 = clip(出队命令 + 平移噪声)；新命令不会在本轮立即执行",(1000,750),31)
    im.save(ASSETS/"fifo_timing.png",dpi=(300,300))

    im,d=canvas(1230)
    labels=[("NORMAL","监测与透传"),("RETRACT","回撤"),("ALIGN","对齐"),("DESCEND","下降"),
            ("CLOSE","闭爪"),("TEST_LIFT","试抬确认"),("LIFT","抬升"),("CARRY","搬运"),
            ("PLACE","下降放置"),("RELEASE","开爪"),("RETREAT","退出"),("NORMAL","回到监测")]
    boxes=[]
    for index,(title,sub) in enumerate(labels):
        row,col=divmod(index,4)
        if row==1: col=3-col
        xy=(75+col*490,100+row*280,475+col*490,270+row*280)
        boxes.append(xy);box(d,xy,title,sub)
    for index in range(11):
        a,b=boxes[index:index+2]
        if index in (3,7): arrow(d,[((a[0]+a[2])/2,a[3]),((b[0]+b[2])/2,b[1])])
        elif index<3 or index>7: arrow(d,[(a[2],(a[1]+a[3])/2),(b[0],(b[1]+b[3])/2)])
        else: arrow(d,[(a[0],(a[1]+a[3])/2),(b[2],(b[1]+b[3])/2)])
    box(d,(170,1000,860,1165),"异常与尝试计数","超时 失持 几何失效",fill="#f6f0e7")
    box(d,(1110,1000,1840,1165),"EXHAUSTED","额度用尽或越界后透传策略",fill="#f6f0e7")
    arrow(d,[(860,1082),(1110,1082)])
    centered(d,"图示为主恢复路径；持物停滞可直接进入LIFT；异常分支与阈值以状态表为准",(1000,920),30)
    im.save(ASSETS/"recovery_state_machine.png",dpi=(300,300))

def assemble():
    sources = [HERE/"sections"/name for name in ("01_04_overview.md","05_06_models.md","07_08_recovery.md","09_12_validation.md")]
    text="\n\n".join(p.read_text(encoding="utf-8").strip() for p in sources)+"\n"
    text=text.replace("../assets/","assets/")
    replacements={
        "BC在两种扰动条件下":"行为克隆（Behavior Cloning，BC）在两种扰动条件下",
        "Full对应结果为":"BC与完整世界模型组合（Full）的对应结果为",
        "DAgger研究讨论了":"数据聚合算法DAgger研究讨论了",
        "PETS说明了":"概率集成与轨迹采样方法PETS说明了",
        "MPC则提供了":"模型预测控制MPC则提供了",
        "系统边界起于HDF5示范":"系统边界起于分层数据格式第五版HDF5示范",
        "设预测步索引为":"递推采用门控循环单元（Gated Recurrent Unit，GRU）。设预测步索引为",
        "延迟队列轨迹审计":"延迟队列轨迹审计",
        "网络由三个宽度为512的全连接隐层构成，每层使用 ReLU 激活。":"网络由三个宽度为512的全连接隐层构成，每层使用修正线性单元ReLU激活。",
        "物体位姿，输出为有限时域":"物体位姿，末端执行器英文缩写为EEF；输出为有限时域",
        "不使用GRU或历史帧队列":"不使用循环网络或历史帧队列",
        "预测差不超过":"预测分差绝对值不超过",
    }
    for old,new in replacements.items(): text=text.replace(old,new,1)
    text=text.replace("系统以专家演示", "系统以专家演示",1)
    text=text.replace("## 5.4 训练目标与参数更新", "![动作块策略结构](assets/policy_network.png)\n\n图5-1 混合密度动作块策略结构\n\n## 5.4 训练目标与参数更新")
    text=text.replace("## 7.3 抓持状态与停滞检测", "![D2执行时序](assets/fifo_timing.png)\n\n图7-1 提交动作与实际执行动作的时序\n\n## 7.3 抓持状态与停滞检测")
    text=text.replace("## 7.5 位姿控制与几何约束", "![恢复主状态路径](assets/recovery_state_machine.png)\n\n图7-2 恢复主路径与异常透传状态\n\n## 7.5 位姿控制与几何约束")
    # Every table receives a caption in both the source and the Word document.
    lines=text.splitlines(); out=[]; chapter=""; counters=defaultdict(int); previous=""
    for index,line in enumerate(lines):
        m=re.match(r"# 第(\d+)章",line)
        if m: chapter=m[1]
        m=re.match(r"# 附录([A-Z])",line)
        if m: chapter=m[1]
        if line.startswith("|") and (index==0 or not lines[index-1].startswith("|")):
            counters[chapter]+=1
            if not re.match(r"表[0-9A-Z]+-\d+",previous):
                labels={"5":["动作块网络张量约定","BC配对训练参数"],"6":["模型分支职责与损失","分支训练数据张量"],"7":["恢复状态与正常转换"],"C":["缩略语与主要符号"]}
                label=labels.get(chapter,["模块参数与接口约定"])
                name=label[min(counters[chapter]-1,len(label)-1)]
                out.extend([f"表{chapter}-{counters[chapter]} {name}",""])
        out.append(line)
        if line.strip(): previous=line.strip()
    text="\n".join(out)+"\n"
    refs=json.loads((HERE/"research/references_verified.json").read_text(encoding="utf-8"))["references"]
    # Number citations by first appearance; match only explicit prose citations,
    # never array indices, intervals or source-code literals.
    order=[]
    pattern=r"(?<=[\u4e00-\u9fff])\[([1-9][0-9]?(?:,[1-9][0-9]?)*)\]"
    for match in re.finditer(pattern,text):
        for old in map(int,match[1].split(',')):
            if old not in order:order.append(old)
    if set(order)!={r['id'] for r in refs}:raise ValueError(f'Citation coverage mismatch: {order}')
    mapping={old:index+1 for index,old in enumerate(order)}
    text=re.sub(pattern,lambda m:'['+','.join(str(mapping[int(old)]) for old in m[1].split(','))+']',text)
    (HERE/'research/document_citation_map.json').write_text(json.dumps(mapping,indent=2)+'\n',encoding='utf-8')
    indexed={r['id']:r for r in refs}
    entries=[]
    for old in order:
        ref=indexed[old];entry=ref['bibliography_zh']
        if ref['stable_url'] not in entry:entry+=' '+ref['stable_url']
        entries.append(f'[{mapping[old]}] {entry}')
    bibliography="# 参考文献\n\n"+"\n\n".join(entries)+"\n\n"
    text=text.replace("# 附录A",bibliography+"# 附录A",1)
    (HERE/"详细设计.md").write_text(text,encoding="utf-8")
    return text,sources


def set_fonts(style, size, east="宋体", west="Times New Roman", bold=False):
    style.font.name=west;style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr()
    rf=rpr.find(qn("w:rFonts"))
    if rf is None: rf=OxmlElement("w:rFonts");rpr.insert(0,rf)
    for name,value in (("ascii",west),("hAnsi",west),("eastAsia",east),("cs",west)):
        rf.set(qn("w:"+name),value)
    for attr in ("asciiTheme","hAnsiTheme","eastAsiaTheme","cstheme"):
        rf.attrib.pop(qn("w:"+attr),None)


def field(paragraph,instruction):
    run=paragraph.add_run()
    begin=OxmlElement("w:fldChar");begin.set(qn("w:fldCharType"),"begin")
    code=OxmlElement("w:instrText");code.set(qn("xml:space"),"preserve");code.text=instruction
    separate=OxmlElement("w:fldChar");separate.set(qn("w:fldCharType"),"separate")
    text=OxmlElement("w:t");text.text=""
    end=OxmlElement("w:fldChar");end.set(qn("w:fldCharType"),"end")
    for element in (begin,code,separate,text,end):run._r.append(element)


class Builder:
    def __init__(self):
        self.doc=Document();self.chapter="";self.math_count=defaultdict(int);self.math_errors=[];self.math_checks=[]
        xsl=Path("C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL")
        self.transform=etree.XSLT(etree.parse(str(xsl)))
        styles=self.doc.styles
        set_fonts(styles["Normal"],12)
        f=styles["Normal"].paragraph_format;f.line_spacing=1.5;f.space_after=Pt(4);f.first_line_indent=Pt(24);f.widow_control=True
        for name,size in (("Title",22),("Heading 1",18),("Heading 2",14),("Heading 3",12)):
            set_fonts(styles[name],size,east="黑体",bold=True)
            pf=styles[name].paragraph_format;pf.first_line_indent=Pt(0);pf.keep_with_next=True
            pf.space_before=Pt(14 if name!="Heading 1" else 6);pf.space_after=Pt(10);pf.line_spacing=1.25
        styles["Heading 1"].paragraph_format.page_break_before=True
        for number in (1,2):
            name=f'toc {number}'
            if name not in styles:styles.add_style(name,1,builtin=True)
            styles[name].element.set(qn('w:styleId'),f'TOC{number}')
            set_fonts(styles[name],11)
            tf=styles[name].paragraph_format;tf.line_spacing=1.1;tf.space_after=Pt(2);tf.space_before=Pt(0);tf.first_line_indent=Pt(0)
            tf.left_indent=Cm(0 if number==1 else .5);tf.keep_with_next=False;tf.page_break_before=False
        styles['TOC Heading'].paragraph_format.page_break_before=True
        styles['TOC Heading'].paragraph_format.keep_with_next=True
        for name,size in (("Caption",10.5),("Body Text",12),("TOC Heading",18)):
            if name in styles: set_fonts(styles[name],size)
        for name in ("Code Design","Equation Design","Reference Design"):
            if name not in styles: styles.add_style(name,1)
        set_fonts(styles["Code Design"],9,east="宋体",west="Consolas")
        cf=styles["Code Design"].paragraph_format;cf.line_spacing=1.15;cf.first_line_indent=Pt(0);cf.left_indent=Cm(.3);cf.space_after=Pt(1)
        set_fonts(styles["Equation Design"],11.5,west="Cambria Math")
        ef=styles["Equation Design"].paragraph_format;ef.first_line_indent=Pt(0);ef.space_after=Pt(9);ef.space_before=Pt(8);ef.line_spacing=1.2
        ef.tab_stops.add_tab_stop(Cm(8),WD_TAB_ALIGNMENT.CENTER);ef.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT)
        set_fonts(styles["Reference Design"],10.5)
        rf=styles["Reference Design"].paragraph_format;rf.first_line_indent=Cm(-.65);rf.left_indent=Cm(.65);rf.line_spacing=1.25;rf.space_after=Pt(8)
        for s in self.doc.sections:
            s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(2.5);s.bottom_margin=Cm(2.4);s.left_margin=Cm(2.7);s.right_margin=Cm(2.3);s.header_distance=Cm(1.2);s.footer_distance=Cm(1.2)
        self.doc.core_properties.title=TITLE
        self.doc.core_properties.subject="详细设计与实现验证"
        self.doc.core_properties.author="项目设计文档"
        self.doc.core_properties.keywords="机器人 抓取搬运 行为克隆 世界模型 延迟 失败恢复"
        self.cover()

    def cover(self):
        p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(70)
        p=self.doc.add_paragraph(TITLE,"Title");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(35)
        for text in ("详细设计说明书","低维状态仿真与受控实验","版本 1.0","实现基线 672fa48","2026年10月22日"):
            p=self.doc.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(14)
        s=self.doc.add_section(WD_SECTION_START.NEW_PAGE)
        s.header.is_linked_to_previous=False;s.footer.is_linked_to_previous=False
        p=s.header.paragraphs[0];p.text=TITLE;p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
        for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
        p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
        field(p,"PAGE")
        num=OxmlElement("w:pgNumType");num.set(qn("w:start"),"1");s._sectPr.append(num)
        settings=self.doc.settings.element
        update=OxmlElement("w:updateFields");update.set(qn("w:val"),"true");settings.append(update)

    def math(self,tex):
        # LaTeX -> MathML -> Office Math is editable, not a raster equation.
        value=tex.strip().replace("\\!","").replace('\\arg\\max','\\operatorname{argmax}').replace('\\argmax','\\operatorname{argmax}')
        mathml=latex_to_mathml(value)
        math_root=etree.fromstring(mathml.encode("utf-8"))
        math_ns='http://www.w3.org/1998/Math/MathML'
        # The Office XSL expects one grouped interior between stretchy fences.
        # Without this normalization a leading unary minus may be discarded.
        for row in list(math_root.iter('{'+math_ns+'}mrow')):
            children=list(row)
            pairs={'(':')','[':']','{':'}','‖':'‖','∥':'∥','|':'|','⌈':'⌉','⌊':'⌋'}
            fenced=(len(children)>2 and children[0].tag=='{'+math_ns+'}mo' and children[-1].tag=='{'+math_ns+'}mo'
                    and pairs.get(children[0].text)==children[-1].text)
            if fenced:
                children[0].set('fence','true');children[0].set('form','prefix')
                children[-1].set('fence','true');children[-1].set('form','postfix')
                inner=etree.Element('{'+math_ns+'}mrow')
                for child in children[1:-1]:row.remove(child);inner.append(child)
                row.insert(1,inner)
        transformed=self.transform(math_root)
        node=transformed.getroot()
        source_text=''.join(math_root.itertext());result_text=''.join(node.itertext())
        if '\\' in result_text or source_text.count('−')!=result_text.count('−'):
            raise ValueError(f'Unconverted command or lost minus in: {tex}')
        self.math_checks.append({'latex':tex,'source_minus_count':source_text.count('−'),'office_minus_count':result_text.count('−'),'unconverted_commands':False})
        if node.tag.endswith("oMathPara"):
            node=node.find("{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath")
        return copy.deepcopy(node)

    def inline(self,p,text):
        pattern=r"(\\\(.*?\\\)|(?<!\$)\$(?!\$).*?(?<!\$)\$(?!\$)|`[^`]+`|\*\*.+?\*\*|\[[^\]]+\]\([^)]+\))"
        for part in re.split(pattern,text):
            if not part:continue
            if part.startswith("\\(") or (part.startswith("$") and part.endswith("$")):
                tex=part[2:-2] if part.startswith("\\(") else part[1:-1]
                try:p._p.append(self.math(tex))
                except Exception as error:self.math_errors.append({"tex":tex,"error":repr(error)});p.add_run(tex)
            elif part.startswith("`"):
                r=p.add_run(part[1:-1]);r.font.name="Consolas";r.font.size=Pt(10)
            elif part.startswith("**"):
                p.add_run(part[2:-2]).bold=True
            elif re.match(r"\[[^\]]+\]\(",part):
                p.add_run(re.match(r"\[([^\]]+)\]",part)[1])
            else:p.add_run(part)

    def table(self,rows):
        table=self.doc.add_table(rows=len(rows),cols=len(rows[0]));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        n=len(rows[0]);weights={2:[.24,.76],3:[.24,.34,.42],4:[.19,.29,.28,.24],5:[.14,.24,.28,.14,.20]}.get(n,[1/n]*n)
        # Numeric result tables need room for compact comparable records.
        if "冻结策略" in rows[0][0]:weights=[.22,.26,.26,.26]
        if rows[0][0]=='接口':weights=[.29,.24,.25,.22]
        for col,w in zip(table.columns,weights):col.width=Cm(16*w)
        pr=table._tbl.tblPr
        borders=OxmlElement("w:tblBorders")
        for edge in ("top","left","bottom","right","insideH","insideV"):
            e=OxmlElement("w:"+edge);e.set(qn("w:val"),"single");e.set(qn("w:sz"),"5");e.set(qn("w:color"),"D9D9D9");borders.append(e)
        pr.append(borders)
        for ri,row in enumerate(rows):
            trpr=table.rows[ri]._tr.get_or_add_trPr();no_split=OxmlElement("w:cantSplit");trpr.append(no_split)
            if ri==0:
                repeat=OxmlElement("w:tblHeader");repeat.set(qn("w:val"),"true");trpr.append(repeat)
            for ci,text in enumerate(row):
                cell=table.cell(ri,ci);cell.width=Cm(16*weights[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr();margins=OxmlElement("w:tcMar")
                for edge,val in (("top",75),("bottom",75),("left",85),("right",85)):
                    e=OxmlElement("w:"+edge);e.set(qn("w:w"),str(val));e.set(qn("w:type"),"dxa");margins.append(e)
                tcpr.append(margins)
                if ri==0:
                    shading=OxmlElement("w:shd");shading.set(qn("w:fill"),"E6ECF0");tcpr.append(shading)
                p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.15
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if len(text)>16 else WD_ALIGN_PARAGRAPH.CENTER
                self.inline(p,text)
                for r in p.runs:r.font.size=Pt(10);r.bold=ri==0
        p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=.5;p.paragraph_format.first_line_indent=Pt(0)

    def render(self,text):
        lines=text.splitlines();i=0;toc=False
        while i<len(lines):
            line=lines[i].strip()
            if not line:i+=1;continue
            if line=="# "+TITLE:i+=1;continue
            if line.startswith("```"):
                i+=1
                while i<len(lines) and not lines[i].startswith("```"):
                    p=self.doc.add_paragraph(lines[i],"Code Design");i+=1
                i+=1;continue
            if line=="$$":
                eq=[];i+=1
                while i<len(lines) and lines[i].strip()!="$$":eq.append(lines[i]);i+=1
                tex="\n".join(eq);self.math_count[self.chapter]+=1
                p=self.doc.add_paragraph(style="Equation Design")
                try:
                    p.add_run("\t");p._p.append(self.math(tex));p.add_run(f"\t({self.chapter}-{self.math_count[self.chapter]})")
                except Exception as error:self.math_errors.append({"tex":tex,"error":repr(error)});p.add_run(tex)
                i+=1;continue
            if line.startswith("|"):
                rows=[]
                while i<len(lines) and lines[i].strip().startswith("|"):
                    cells=[c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-+:?",c) for c in cells):rows.append(cells)
                    i+=1
                self.table(rows);continue
            if re.match(r"!\[.*?\]\(",line):
                match=re.match(r"!\[(.*?)\]\((.*?)\)",line);path=HERE/match[2]
                p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
                p.add_run().add_picture(str(path),width=Cm(15.7));i+=1;continue
            if line.startswith("#"):
                level=len(line)-len(line.lstrip("#"));title=line[level:].strip()
                if title.startswith("第1章") and not toc:
                    p=self.doc.add_paragraph("目录","TOC Heading");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0)
                    field(p,'TOC \\o "1-2" \\h \\z \\u');toc=True
                chapter=re.match(r"第(\d+)章",title)
                if chapter:self.chapter=chapter[1]
                actual=1 if title in ("摘要","Abstract") else min(level,3)
                p=self.doc.add_paragraph(title,"Heading "+str(actual));p.alignment=WD_ALIGN_PARAGRAPH.CENTER if actual==1 else WD_ALIGN_PARAGRAPH.LEFT
                i+=1;continue
            if re.match(r"[表图][0-9A-Z]+-\d+\s",line):
                p=self.doc.add_paragraph(line,"Caption");p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(6);p.paragraph_format.space_before=Pt(6);p.paragraph_format.keep_with_next=line.startswith("表")
                i+=1;continue
            parts=[line];i+=1
            while i<len(lines) and lines[i].strip() and not re.match(r"(#|\||```|\$\$|!\[|[表图][0-9A-Z]+-\d+)",lines[i].strip()):
                parts.append(lines[i].strip());i+=1
            value=" ".join(parts)
            p=self.doc.add_paragraph(style="Reference Design" if re.match(r"\[\d+\] ",value) else "Normal")
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if p.style.name=='Reference Design' or value.startswith('表中未给出目录前缀') else WD_ALIGN_PARAGRAPH.JUSTIFY
            self.inline(p,value)
        if self.math_errors:raise RuntimeError(json.dumps(self.math_errors,ensure_ascii=False))
        return self.doc


def main():
    parser=argparse.ArgumentParser();parser.add_argument("--replace-generated",action="store_true");args=parser.parse_args()
    output=HERE/(TITLE+".docx")
    if output.exists() and not args.replace_generated:raise FileExistsError(output)
    save_diagrams()
    from plot_for_document import generate
    generate()
    text,sources=assemble()
    builder=Builder();document=builder.render(text);document.save(output)
    source_paths=[*sources,HERE/"research/references_verified.json",ROOT/"common.py",ROOT/"world_model.py",*(ROOT/"robot_can_v2").glob("*.py"),RUN/"REPORT.md",RUN/"analysis/confirmation/summary.json",RUN/"analysis/confirmation/cells.csv",RUN/"analysis/confirmation/paired_effects.csv"]
    stats={"implementation_baseline":"672fa48","document_kind":"detailed technical design","chapters":12,
           "cjk_characters":len(re.findall(r"[\u4e00-\u9fff]",text)),"characters":len(text),
           "equations":sum(builder.math_count.values()),"inline_math_failures":len(builder.math_errors),
           "tables":len(document.tables),"figures":len(document.inline_shapes),"references":10,
           "math_pipeline":"latex2mathml 3.81.1 and Microsoft Office MML2OMML XSL; editable Office Math",
           "sources_sha256":{p.relative_to(ROOT).as_posix():sha(p) for p in source_paths},
           "outputs_sha256":{p.name:sha(p) for p in (output,HERE/"详细设计.md")},
           "note":"Build completion is not visual QA; Word field update and every-page rendering inspection follow separately."}
    (HERE/"build_manifest.json").write_text(json.dumps(stats,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    (HERE/'research/math_conversion_checks.json').write_text(json.dumps({'scope':'Unconverted command and minus preservation checks; visual inspection remains required.','checks':builder.math_checks},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({k:v for k,v in stats.items() if k not in ("sources_sha256","outputs_sha256")},ensure_ascii=False))


if __name__=="__main__":main()
